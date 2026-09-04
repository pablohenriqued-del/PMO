#!/usr/bin/env python3
"""Gera a apresentação comercial PRISMA PMO em PPTX usando telas reais."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image
import os

SCR = "/app/presentation_assets/screens"
LOGO = "/app/presentation_assets/prisma_logo.jpeg"
OUT = "/app/presentation_assets/PRISMA_PMO_Apresentacao.pptx"

# Paleta
BLACK = RGBColor(0x0A, 0x0A, 0x0A)
PANEL = RGBColor(0x15, 0x15, 0x17)
CARD = RGBColor(0x1C, 0x1C, 0x20)
RED = RGBColor(0xE5, 0x09, 0x14)
DARKRED = RGBColor(0xB2, 0x07, 0x10)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0xA0, 0xA0, 0xA8)
LGRAY = RGBColor(0xC8, 0xC8, 0xD0)
GREEN = RGBColor(0x2E, 0xCC, 0x71)
GOLD = RGBColor(0xF5, 0xB3, 0x41)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    bg.fill.solid(); bg.fill.fore_color.rgb = BLACK
    bg.line.fill.background()
    bg.shadow.inherit = False
    return s


def rect(s, x, y, w, h, color, line=None, line_w=None, shape=MSO_SHAPE.RECTANGLE):
    sp = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if color is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = color
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(line_w or 1)
    sp.shadow.inherit = False
    return sp


def txt(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, sp_after=6, line_sp=1.05):
    """runs: list of paragraphs; each paragraph = list of (text,size,color,bold) tuples."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    from pptx.oxml.ns import qn
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(sp_after)
        p.line_spacing = line_sp
        for (t, sz, col, bold) in para:
            r = p.add_run(); r.text = t
            r.font.size = Pt(sz); r.font.color.rgb = col; r.font.bold = bold
            r.font.name = "Arial"
    return tb


def img_fit(s, path, x, y, maxw, maxh):
    im = Image.open(path); iw, ih = im.size
    ar = iw / ih
    w = maxw; h = w / ar
    if h > maxh:
        h = maxh; w = h * ar
    px = x + (maxw - w) / 2
    py = y + (maxh - h) / 2
    # borda
    rect(s, px - 0.03, py - 0.03, w + 0.06, h + 0.06, None, line=RED, line_w=1.5)
    s.shapes.add_picture(path, Inches(px), Inches(py), Inches(w), Inches(h))
    return px, py, w, h


def kicker_bar(s, x, y):
    rect(s, x, y, 0.5, 0.09, RED)


# ---------------- 1. CAPA ----------------
s = slide()
rect(s, 0, 0, 13.333, 7.5, BLACK)
# faixa radial simulada
rect(s, 0, 0, 13.333, 2.4, PANEL)
rect(s, 0, 2.38, 13.333, 0.04, RED)
s.shapes.add_picture(LOGO, Inches(0.85), Inches(0.7), Inches(1.7), Inches(1.7))
txt(s, 2.8, 0.75, 9.5, 1.7, [
    [("PRISMA PMO", 46, WHITE, True)],
    [("Plataforma Executiva de Gestão de Portfólio & Projetos", 18, RED, True)],
], anchor=MSO_ANCHOR.MIDDLE)
txt(s, 0.9, 3.1, 11.5, 2.4, [
    [("Uma plataforma corporativa completa para PMOs modernos:", 20, LGRAY, False)],
    [("delivery, governança, analytics, inteligência artificial e CRM —", 20, LGRAY, False)],
    [("tudo em um único ambiente, com design premium e pronto para escalar.", 20, LGRAY, False)],
], sp_after=8)
txt(s, 0.9, 6.4, 11.5, 0.8, [
    [("Apresentação Executiva & Comercial", 14, GRAY, True), ("   •   White-label / Marca própria", 14, RED, True)],
])

# ---------------- 2. O QUE É ----------------
s = slide()
kicker_bar(s, 0.9, 0.75)
txt(s, 0.9, 0.9, 11.5, 1.0, [[("O que é o PRISMA PMO", 34, WHITE, True)]])
txt(s, 0.9, 1.8, 11.6, 1.0, [[(
    "Um cockpit único que conecta estratégia e execução. Do planejamento ao encerramento, o PRISMA "
    "centraliza projetos, orçamento, riscos, capacidade e demandas — com inteligência artificial que "
    "antecipa gargalos e apoia a tomada de decisão executiva.", 16, LGRAY, False)]])
cards = [
    ("Centralização", "Portfólio inteiro em uma só fonte de verdade"),
    ("Inteligência", "IA que prioriza, alerta e recomenda ações"),
    ("Governança", "Riscos, causas-raiz e lições aprendidas"),
    ("Visão Executiva", "Dashboards e relatórios prontos para a diretoria"),
]
cx = 0.9; cw = 2.85; gap = 0.23; cy = 3.4; ch = 3.2
for i, (t, d) in enumerate(cards):
    x = cx + i * (cw + gap)
    rect(s, x, cy, cw, ch, CARD, line=RGBColor(0x33,0x33,0x38), line_w=0.75)
    rect(s, x, cy, cw, 0.09, RED)
    txt(s, x+0.25, cy+0.4, cw-0.5, 0.7, [[(t, 18, WHITE, True)]])
    txt(s, x+0.25, cy+1.3, cw-0.5, 1.6, [[(d, 14, GRAY, False)]])

# ---------------- 3. O DESAFIO ----------------
s = slide()
kicker_bar(s, 0.9, 0.75)
txt(s, 0.9, 0.9, 11.5, 1.0, [[("O desafio dos PMOs hoje", 34, WHITE, True)]])
pains = [
    ("Dados fragmentados", "Planilhas soltas, ferramentas desconectadas e retrabalho constante."),
    ("Reporting manual", "Horas para montar status reports que ficam desatualizados no dia seguinte."),
    ("Riscos invisíveis", "Gargalos e atrasos só aparecem quando já viraram crise."),
    ("Baixa previsibilidade", "Dificuldade de enxergar capacidade, ROI e saúde real do portfólio."),
]
cy = 2.1
for i, (t, d) in enumerate(pains):
    y = cy + i * 1.18
    rect(s, 0.9, y, 11.5, 1.0, PANEL)
    rect(s, 0.9, y, 0.09, 1.0, RED)
    txt(s, 1.25, y+0.12, 3.6, 0.8, [[(t, 17, WHITE, True)]], anchor=MSO_ANCHOR.MIDDLE)
    txt(s, 5.0, y+0.12, 7.2, 0.8, [[(d, 15, GRAY, False)]], anchor=MSO_ANCHOR.MIDDLE)

# ---------------- 4. A SOLUÇÃO (6 MÓDULOS) ----------------
s = slide()
kicker_bar(s, 0.9, 0.75)
txt(s, 0.9, 0.9, 11.5, 1.0, [[("A solução: 6 módulos integrados", 34, WHITE, True)]])
mods = [
    ("Dashboard Executivo", "Visão 360° do portfólio em tempo real"),
    ("Delivery", "Projetos, Timeline, Budget e Capacity"),
    ("Analytics & IA", "AI Copilot, Status Report e Regional"),
    ("Governança PMO", "Riscos, RCA, Lições e Playbook"),
    ("CRM & Demandas", "Pipeline Kanban com conversão em projeto"),
    ("Administração", "Usuários, permissões e integrações"),
]
cw = 3.7; ch = 1.85; gapx = 0.3; gapy = 0.35; sx = 0.9; sy = 2.15
for i, (t, d) in enumerate(mods):
    col = i % 3; row = i // 3
    x = sx + col * (cw + gapx); y = sy + row * (ch + gapy)
    rect(s, x, y, cw, ch, CARD, line=RGBColor(0x33,0x33,0x38), line_w=0.75)
    rect(s, x, y, 0.09, ch, RED)
    txt(s, x+0.3, y+0.28, cw-0.55, 0.6, [[(t, 17, WHITE, True)]])
    txt(s, x+0.3, y+0.95, cw-0.55, 0.8, [[(d, 13.5, GRAY, False)]])

# ---------------- 5. DIFERENCIAIS ----------------
s = slide()
kicker_bar(s, 0.9, 0.75)
txt(s, 0.9, 0.9, 11.5, 1.0, [[("Por que o PRISMA PMO", 34, WHITE, True)]])
diffs = [
    ("IA Autônoma (Agentic)", "Copiloto que monitora, alerta e sugere realocações automaticamente."),
    ("White-label", "Marca própria, cores e domínio — pronto para revenda."),
    ("Multilíngue PT / EN / ES", "Times globais no mesmo ambiente."),
    ("Exportação PDF / PPT / PNG", "Relatórios executivos com um clique."),
    ("SSO Corporativo", "Login sem fricção via sessão corporativa."),
    ("Design Premium Dark", "Experiência moderna, rápida e responsiva."),
]
cw = 5.7; ch = 1.35; gapx = 0.25; gapy = 0.25; sx = 0.9; sy = 2.1
for i, (t, d) in enumerate(diffs):
    col = i % 2; row = i // 2
    x = sx + col * (cw + gapx); y = sy + row * (ch + gapy)
    rect(s, x, y, cw, ch, PANEL, line=RGBColor(0x33,0x33,0x38), line_w=0.75)
    rect(s, x, y, 0.09, ch, RED)
    txt(s, x+0.3, y+0.16, cw-0.55, 0.5, [[(t, 16, WHITE, True)]])
    txt(s, x+0.3, y+0.66, cw-0.55, 0.6, [[(d, 13, GRAY, False)]])

# ---------------- WALKTHROUGH ----------------
features = [
    ("00_login", "ACESSO", "Login & SSO Corporativo",
     "Autenticação segura sem fricção para toda a organização.",
     ["Login SSO corporativo sem senha", "Acesso externo com e-mail e senha", "Tokens JWT + cookies httpOnly"]),
    ("01_dashboard", "VISÃO EXECUTIVA", "Dashboard Executivo",
     "O portfólio inteiro em uma tela, atualizado em tempo real.",
     ["Saúde do portfólio e ROI global", "Evolução financeira e top projetos", "Atividades recentes relevantes"]),
    ("02_projects", "DELIVERY", "Gestão de Projetos & Portfólio",
     "Do planejamento à entrega, com múltiplas visões.",
     ["Views Grid, Kanban, Tabela e Gantt", "Filtros por país, área, tipo e prioridade", "Milestones, tarefas, budget e documentos"]),
    ("03_timeline", "DELIVERY", "Timeline & Gantt Executivo",
     "Roadmap anual elegante para apresentar à diretoria.",
     ["Trimestres Q1–Q4 em uma linha do tempo", "Key Milestones em destaque (losangos)", "Exportação em PNG para slides"]),
    ("04_budget", "DELIVERY", "Orçamento & Financeiro",
     "Controle financeiro de ponta a ponta do portfólio.",
     ["Orçado vs. realizado por projeto", "Receita esperada vs. gerada", "Relatórios financeiros exportáveis"]),
    ("05_capacity", "DELIVERY", "Capacity Planning & Heatmap",
     "Previsão de carga e prevenção de burnout das equipes.",
     ["Capacidade calculada por pessoa", "Mapa de calor de alocação", "Alertas de sobrecarga (>100%)"]),
    ("06_ai_copilot", "INTELIGÊNCIA", "PMO AI Copilot",
     "O cérebro autônomo que antecipa problemas e recomenda ações.",
     ["Priorização orientada por tendências", "Detecção de gargalos e riscos de atraso", "Realocação de recursos em 1 clique"]),
    ("07_status_report", "INTELIGÊNCIA", "Status Report Executivo",
     "Relatório de status vivo, gerado automaticamente.",
     ["Indicadores de saúde (cronograma, budget, escopo)", "Riscos críticos vinculados ao projeto", "Exportação em PDF / PPT"]),
    ("08_regional", "ANALYTICS", "Regional Analytics",
     "Performance financeira comparada por país e mercado.",
     ["Receita, lucratividade e ROI por região", "Eficiência operacional por país", "Comparativo entre mercados"]),
    ("09_innovation", "ANALYTICS", "Innovation Radar",
     "Mapa de tecnologias e oportunidades de inovação.",
     ["Quadrantes Adopt / Trial / Assess / Hold", "Oportunidades de alto impacto", "Portfólio de inovação estruturado"]),
    ("10_risk_radar", "GOVERNANÇA", "Risk Radar",
     "Gestão de riscos conectada diretamente aos projetos.",
     ["Riscos vinculados a cada projeto", "Probabilidade x impacto", "Planos de mitigação e monitoramento"]),
    ("11_root_cause", "GOVERNANÇA", "Root Cause Analysis",
     "Investigação sistemática de incidentes e falhas.",
     ["Identificação de causas-raiz", "Ações corretivas rastreáveis", "Impacto quantificado"]),
    ("12_lessons", "GOVERNANÇA", "Lessons Learned",
     "Base de conhecimento para melhoria contínua.",
     ["Aprendizados capturados por projeto", "Reutilização entre iniciativas", "Cultura de melhoria contínua"]),
    ("13_playbook", "GOVERNANÇA", "PMO Playbook",
     "Framework de entrega padronizado da organização.",
     ["Metodologias ágeis + governança", "Templates e boas práticas", "Onboarding acelerado de times"]),
    ("14_crm", "COMERCIAL", "CRM & Demandas (Kanban)",
     "Pipeline visual de leads e demandas do PMO.",
     ["Kanban com arrastar e soltar", "Etapas do funil configuráveis", "Conversão automática em projeto (Closed Won)"]),
    ("15_admin", "ADMINISTRAÇÃO", "Gestão de Usuários",
     "Controle total de acessos e permissões.",
     ["CRUD de usuários e perfis", "Administradores de portal", "Integração LDAP / Active Directory"]),
    ("16_features", "PLATAFORMA", "Guia & Funcionalidades",
     "Toda a plataforma documentada para o usuário.",
     ["Regras de negócio explicadas", "Recursos premium detalhados", "Onboarding self-service"]),
]

for fn, kick, title, desc, bullets in features:
    s = slide()
    # painel esquerdo texto
    rect(s, 0, 0, 4.55, 7.5, PANEL)
    rect(s, 4.55, 0, 0.05, 7.5, RED)
    kicker_bar(s, 0.6, 0.7)
    txt(s, 0.6, 0.82, 3.6, 0.5, [[(kick, 13, RED, True)]])
    txt(s, 0.55, 1.25, 3.7, 1.8, [[(title, 27, WHITE, True)]])
    txt(s, 0.6, 3.0, 3.7, 1.0, [[(desc, 14.5, LGRAY, False)]])
    # bullets
    by = 4.2
    for b in bullets:
        rect(s, 0.62, by+0.12, 0.14, 0.14, RED, shape=MSO_SHAPE.OVAL)
        txt(s, 0.92, by, 3.4, 0.7, [[(b, 13.5, GRAY, False)]])
        by += 0.85
    # screenshot
    img_fit(s, os.path.join(SCR, fn + ".jpeg"), 4.95, 0.9, 7.9, 5.7)
    # rodapé
    txt(s, 4.95, 6.75, 7.9, 0.4, [[("PRISMA PMO", 10, GRAY, True), ("  —  tela real da plataforma", 10, RGBColor(0x66,0x66,0x6c), False)]])

# ---------------- BENEFÍCIOS / ROI ----------------
s = slide()
kicker_bar(s, 0.9, 0.75)
txt(s, 0.9, 0.9, 11.5, 1.0, [[("Benefícios & impacto", 34, WHITE, True)]])
kpis = [
    ("+30%", "previsibilidade de entregas"),
    ("-25%", "tempo em reporting manual"),
    ("1", "fonte única de verdade do portfólio"),
    ("100%", "decisões orientadas por dados"),
]
cw = 2.85; gap = 0.23; sy = 2.2; ch = 2.2; sx = 0.9
for i, (v, d) in enumerate(kpis):
    x = sx + i * (cw + gap)
    rect(s, x, sy, cw, ch, CARD, line=RGBColor(0x33,0x33,0x38), line_w=0.75)
    txt(s, x, sy+0.35, cw, 1.0, [[(v, 40, RED, True)]], align=PP_ALIGN.CENTER)
    txt(s, x+0.2, sy+1.4, cw-0.4, 0.7, [[(d, 13.5, GRAY, False)]], align=PP_ALIGN.CENTER)
txt(s, 0.9, 4.9, 11.5, 1.4, [
    [("Menos esforço operacional, mais foco em estratégia.", 18, WHITE, True)],
    [("O PRISMA transforma dados dispersos em decisões executivas rápidas — reduzindo custos de "
      "coordenação e aumentando a taxa de sucesso dos projetos.", 15, GRAY, False)],
], sp_after=10)
txt(s, 0.9, 6.7, 11.5, 0.4, [[("* Estimativas ilustrativas — ajustáveis conforme o cliente.", 11, RGBColor(0x66,0x66,0x6c), False)]])

# ---------------- SEGURANÇA & INTEGRAÇÕES ----------------
s = slide()
kicker_bar(s, 0.9, 0.75)
txt(s, 0.9, 0.9, 11.5, 1.0, [[("Segurança, SSO & integrações", 34, WHITE, True)]])
items = [
    ("SSO Corporativo", "Login via sessão corporativa, sem senhas expostas."),
    ("JWT + Cookies httpOnly", "Sessões seguras com tokens de acesso e refresh."),
    ("LDAP / Active Directory", "Sincronização de usuários e alocação de recursos."),
    ("Notificações & E-mail", "Alertas automáticos de alocação e riscos."),
    ("Importação Planner / Monday / CSV", "Migração de cronogramas e demandas."),
    ("Exportação PDF / PPT / PNG", "Relatórios e apresentações executivas."),
]
cw = 5.7; ch = 1.35; gapx = 0.25; gapy = 0.25; sx = 0.9; sy = 2.1
for i, (t, d) in enumerate(items):
    col = i % 2; row = i // 2
    x = sx + col * (cw + gapx); y = sy + row * (ch + gapy)
    rect(s, x, y, cw, ch, PANEL, line=RGBColor(0x33,0x33,0x38), line_w=0.75)
    rect(s, x, y, 0.09, ch, RED)
    txt(s, x+0.3, y+0.16, cw-0.55, 0.5, [[(t, 15.5, WHITE, True)]])
    txt(s, x+0.3, y+0.66, cw-0.55, 0.6, [[(d, 12.5, GRAY, False)]])

# ---------------- PLANOS & PREÇOS ----------------
s = slide()
kicker_bar(s, 0.9, 0.75)
txt(s, 0.9, 0.9, 11.5, 1.0, [[("Planos & preços", 34, WHITE, True)]])
txt(s, 0.9, 1.75, 11.5, 0.5, [[("Escolha o plano que acompanha o crescimento da sua operação.", 15, GRAY, False)]])
tiers = [
    ("STARTER", "R$ 1.900", "/mês", False,
     ["Até 10 usuários", "Dashboard Executivo", "Projetos, Timeline & Budget", "Suporte por e-mail"]),
    ("PROFESSIONAL", "R$ 4.900", "/mês", True,
     ["Até 50 usuários", "Tudo do Starter +", "Analytics, AI Copilot & CRM", "Risk Radar & Status Report", "Suporte prioritário"]),
    ("ENTERPRISE", "Sob consulta", "", False,
     ["Usuários ilimitados", "SSO & White-label", "Integrações dedicadas", "Onboarding e SLA dedicados", "Gerente de sucesso"]),
]
cw = 3.75; gap = 0.35; sx = 0.9; sy = 2.4; ch = 4.65
for i, (name, price, per, hl, feats) in enumerate(tiers):
    x = sx + i * (cw + gap)
    base = CARD if not hl else RGBColor(0x22, 0x0a, 0x0c)
    ln = RGBColor(0x33,0x33,0x38) if not hl else RED
    rect(s, x, sy, cw, ch, base, line=ln, line_w=1.5 if hl else 0.75)
    if hl:
        rect(s, x, sy, cw, 0.45, RED)
        txt(s, x, sy+0.02, cw, 0.42, [[("MAIS POPULAR", 11, WHITE, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    off = 0.45 if hl else 0.0
    txt(s, x+0.3, sy+0.35+off, cw-0.6, 0.5, [[(name, 15, RED if not hl else WHITE, True)]])
    txt(s, x+0.3, sy+0.95+off, cw-0.6, 0.8, [[(price, 30, WHITE, True), (per, 13, GRAY, False)]])
    fy = sy + 1.9 + off
    for f in feats:
        txt(s, x+0.3, fy, cw-0.6, 0.5, [[("•  ", 13, RED, True), (f, 12.5, LGRAY, False)]])
        fy += 0.4
txt(s, 0.9, 7.15, 11.5, 0.4, [[("* Valores ilustrativos para referência comercial — personalizáveis por contrato.", 11, RGBColor(0x66,0x66,0x6c), False)]])

# ---------------- CTA ----------------
s = slide()
rect(s, 0, 0, 13.333, 7.5, BLACK)
rect(s, 0, 3.0, 13.333, 0.05, RED)
s.shapes.add_picture(LOGO, Inches(5.82), Inches(0.9), Inches(1.7), Inches(1.7))
txt(s, 0.9, 3.2, 11.5, 1.2, [[("Vamos transformar a gestão do seu portfólio?", 32, WHITE, True)]], align=PP_ALIGN.CENTER)
txt(s, 0.9, 4.5, 11.5, 1.0, [
    [("Agende uma demonstração personalizada do PRISMA PMO", 18, LGRAY, False)],
    [("e veja a plataforma funcionando com os dados da sua empresa.", 18, LGRAY, False)],
], align=PP_ALIGN.CENTER, sp_after=6)
btn = rect(s, 5.17, 6.0, 3.0, 0.75, RED, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
txt(s, 5.17, 6.02, 3.0, 0.7, [[("Solicitar demonstração", 15, WHITE, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

prs.save(OUT)
print("SAVED", OUT, "slides:", len(prs.slides._sldIdLst))
